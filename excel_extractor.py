# -*- coding: utf-8 -*-
"""
엑셀 값 추출 (F6)
- 엑셀의 '항목명' 칸과 그 옆 '값' 칸을 쌍으로 읽어 딕셔너리로 뽑아낸다.
- 추출값은 원본과 100% 일치해야 하므로 AI를 쓰지 않고 openpyxl로 직접 읽는다 (오차 없음).
- 문장으로 바꾸는 건 ai_writer.to_gaejoshik(F1)에 넘겨서 재사용한다.
"""
import openpyxl


def extract_label_value_pairs(path: str, sheet_name: str = None) -> dict:
    """엑셀에서 라벨 셀과 그 바로 오른쪽 셀(값)을 쌍으로 뽑아 딕셔너리로 반환한다."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active

    result = {}
    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str) or not cell.value.strip():
                continue
            label = cell.value.strip().rstrip(":：")
            next_cell = ws.cell(row=cell.row, column=cell.column + 1)
            if next_cell.value is not None:
                result[label] = next_cell.value
    return result


def _format_value(v):
    """숫자는 천단위 쉼표를 넣어 읡기 쉽게 만든다 (원본 숫자 자체는 바뀌지 않음)."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return f"{v:,}"
    return v


def values_to_draft_text(values: dict) -> str:
    """추출된 값들을 자연어 초안 문장으로 엮는다 (F1 입력용)."""
    parts = [f"{k}은(는) {_format_value(v)}" for k, v in values.items()]
    return ", ".join(parts) + " 임을 자연스러운 보고서 문장으로 정리해줘 (단위를 문맥에 맞게 붙여줘)"


if __name__ == "__main__":
    # 테스트용 샘플 엑셀 생성
    sample_path = "_테스트실적표.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "참가인원"
    ws["B1"] = 342
    ws["A2"] = "전월대비 증가율"
    ws["B2"] = "12%"
    ws["A3"] = "예산 집행액"
    ws["B3"] = 1850000
    wb.save(sample_path)

    values = extract_label_value_pairs(sample_path)
    print("추출된 값:", values)
    assert values["참가인원"] == 342, "원본과 불일치!"
    assert values["예산 집행액"] == 1850000, "원본과 불일치!"
    print("검증: 원본 숫자와 100% 일치")

    draft = values_to_draft_text(values)
    print("자연어 초안:", draft)

    from ai_writer import to_gaejoshik
    print("공문서 개조식:", to_gaejoshik(draft))

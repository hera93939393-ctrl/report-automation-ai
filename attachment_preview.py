"""attachment_preview.py — "+"로 첨부한 문서를 첨부 즉시 작은 미리보기로
보여주기 위한 도구. HWP/HWPX는 이미지 미리보기(1페이지만), 엑셀/PDF는
텍스트 미리보기(상위 몇 줄)를 만든다. source_reader.py의
read_hwp_source_isolated()와 같은 이유로, HWP 미리보기는 반드시 별도
프로세스에서 만든다(같은 프로세스 안의 다른 Hwp 인스턴스가 채팅 세션이
이미 열어둔 HwpReport의 COM 연결을 깨뜨리는 pyhwpx의 한계, source_reader.py
주석 참고).

실측 확인(2026-09-07): create_page_image(path, pgno=1, resolution=60,
format="gif")로 만든 1페이지 미리보기는 문서 한 장 기준 수 KB 수준이다
(직접 측정: 짧은 문서 1장 gif@60dpi 약 6.8KB) - bmp(기본 포맷)는 무압축이라
96dpi 한 장에도 3.5MB나 나가 첨부할 때마다 용량이 급격히 쌓이므로, 이
도구는 반드시 gif+저해상도 조합을 쓴다(사용자가 명시적으로 우려한 용량
관리 요구사항)."""
import glob
import os
import subprocess
import sys
import tempfile

_PREVIEW_DIR = os.path.join(tempfile.gettempdir(), "_hwp_attachment_previews")
_MAX_PREVIEWS = 5
_PREVIEW_RESOLUTION = 60


def generate_text_preview(path: str, max_lines: int = 5) -> str:
    """엑셀/PDF의 상위 max_lines줄만 텍스트로 미리보기를 만든다. 이미지이
    아니라 문자열을 그대로 반환한다(채팅 말풍선에 바로 표시할 수 있게).
    지원하지 않는 형식이거나 읽기 실패하면 빈 문자열(예외 없음, 이 모듈의
    다른 함수들과 같은 관례)."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True)
            ws = wb.active
            lines = []
            for row in ws.iter_rows(max_row=max_lines, values_only=True):
                lines.append(", ".join(str(v) for v in row if v is not None))
            return "\n".join(lines)
        if ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                text = pdf.pages[0].extract_text() or ""
            return "\n".join(text.split("\n")[:max_lines])
    except Exception:
        return ""
    return ""


def _generate_hwp_preview(path: str, out_path: str) -> bool:
    """path(.hwp/.hwpx)의 1페이지를 out_path(gif)로 저장한다. 이 함수 자체는
    현재 프로세스에서 새 Hwp 인스턴스를 만들므로, 채팅 세션이 이미 다른
    Hwp를 열어둔 상태에서 절대 직접 호출하면 안 된다 - 반드시
    generate_hwp_preview_isolated()를 통해 별도 프로세스에서만 실행한다."""
    from pyhwpx import Hwp
    hwp = None
    try:
        hwp = Hwp(visible=False, new=True)
        if not hwp.open(path):
            return False
        return bool(hwp.create_page_image(out_path, pgno=1, resolution=_PREVIEW_RESOLUTION, format="gif"))
    except Exception:
        return False
    finally:
        if hwp is not None:
            hwp.quit()


def generate_hwp_preview_isolated(path: str) -> str | None:
    """_generate_hwp_preview()를 별도 프로세스에서 실행해, 성공하면 저장된
    미리보기 이미지 경로를, 실패하면 None을 반환한다. source_reader.py의
    read_hwp_source_isolated()와 동일한 subprocess 격리 패턴. 호출 성공
    시마다 _prune_old_previews()로 오래된 미리보기를 정리한다."""
    os.makedirs(_PREVIEW_DIR, exist_ok=True)
    out_path = os.path.join(_PREVIEW_DIR, f"{os.path.basename(path)}_{os.getpid()}.gif")
    module_dir = os.path.dirname(os.path.abspath(__file__))
    script = (
        "import sys; sys.path.insert(0, sys.argv[3]); "
        "from attachment_preview import _generate_hwp_preview; "
        "ok = _generate_hwp_preview(sys.argv[1], sys.argv[2]); "
        "sys.exit(0 if ok else 1)"
    )
    try:
        result = subprocess.run(
            [sys.executable, "-c", script, path, out_path, module_dir],
            capture_output=True, text=True, timeout=60,
        )
        if result.returncode != 0 or not os.path.exists(out_path):
            return None
        _prune_old_previews()
        return out_path
    except Exception:
        return None


def _prune_old_previews(keep: int = _MAX_PREVIEWS) -> None:
    """_PREVIEW_DIR에 쌓인 미리보기 이미지 중 오래된 것부터 지워 최근
    keep개만 남긴다(용량 관리 요구사항). 파일 개수가 keep 이하면 아무 것도
    안 한다."""
    if not os.path.isdir(_PREVIEW_DIR):
        return
    files = sorted(glob.glob(os.path.join(_PREVIEW_DIR, "*.gif")), key=os.path.getmtime)
    for old_file in files[:-keep] if len(files) > keep else []:
        try:
            os.remove(old_file)
        except OSError:
            pass


def _selftest_generate_hwp_preview_isolated_creates_small_file():
    """미리보기 이미지가 실제로 생성되고, gif+저해상도 조합 덕분에 충분히
    작은지(용량 관리 요구사항) 확인한다."""
    import time
    from pyhwpx import Hwp

    path = os.path.join(tempfile.gettempdir(), "_test_미리보기_원본.hwp")
    setup = Hwp(visible=False, new=True)
    setup.insert_text("미리보기 테스트 문서입니다.")
    setup.save_as(path)
    setup.quit()
    time.sleep(2)

    try:
        preview_path = generate_hwp_preview_isolated(path)
        assert preview_path is not None
        assert os.path.exists(preview_path)
        size = os.path.getsize(preview_path)
        assert size < 200_000, f"미리보기 용량이 예상보다 큼: {size} bytes"
        print("generate_hwp_preview_isolated 통과: 용량", size, "bytes")
    finally:
        os.remove(path)
        for f in glob.glob(os.path.join(_PREVIEW_DIR, "*.gif")):
            os.remove(f)


def _selftest_generate_hwp_preview_isolated_missing_file_returns_none():
    result = generate_hwp_preview_isolated("존재하지_않는_미리보기.hwp")
    assert result is None, result
    print("generate_hwp_preview_isolated(파일없음) 통과")


def _selftest_generate_text_preview_excel():
    import openpyxl
    path = os.path.join(tempfile.gettempdir(), "_test_미리보기.xlsx")
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["항목", "금액"])
    ws.append(["인건비", 1000000])
    wb.save(path)
    try:
        preview = generate_text_preview(path, max_lines=5)
        assert "항목" in preview and "인건비" in preview, preview
        print("generate_text_preview(엑셀) 통과:", repr(preview))
    finally:
        os.remove(path)


def _selftest_generate_text_preview_pdf():
    from fpdf import FPDF
    path = os.path.join(tempfile.gettempdir(), "_test_미리보기.pdf")
    pdf = FPDF()
    pdf.add_font("Malgun", fname="C:/Windows/Fonts/malgun.ttf")
    pdf.add_page(); pdf.set_font("Malgun", size=12)
    pdf.cell(0, 10, "예산은 1,850,000원입니다")
    pdf.output(path)
    try:
        preview = generate_text_preview(path, max_lines=5)
        assert "1,850,000" in preview or "1850000" in preview, preview
        print("generate_text_preview(PDF) 통과:", repr(preview))
    finally:
        os.remove(path)


def _selftest_prune_old_previews_keeps_recent_n():
    """미리보기가 keep개보다 많이 쌓이면 오래된 것부터 지워 최근 keep개만
    남아야 한다."""
    os.makedirs(_PREVIEW_DIR, exist_ok=True)
    test_files = []
    for i in range(8):
        p = os.path.join(_PREVIEW_DIR, f"_test_prune_{i}.gif")
        with open(p, "w") as f:
            f.write("dummy")
        os.utime(p, (i, i))  # 오래된 순서대로 mtime을 강제로 다르게 설정
        test_files.append(p)
    try:
        _prune_old_previews(keep=5)
        remaining = sorted(glob.glob(os.path.join(_PREVIEW_DIR, "_test_prune_*.gif")))
        assert len(remaining) == 5, remaining
        # 가장 오래된 3개(_0,_1,_2)는 지워지고 최근 5개(_3~_7)만 남아야 함
        for i in range(3):
            assert os.path.join(_PREVIEW_DIR, f"_test_prune_{i}.gif") not in remaining
        print("_prune_old_previews(최근 5개만 유지) 통과")
    finally:
        for p in test_files:
            if os.path.exists(p):
                os.remove(p)


if __name__ == "__main__":
    _selftest_generate_hwp_preview_isolated_creates_small_file()
    _selftest_generate_hwp_preview_isolated_missing_file_returns_none()
    _selftest_generate_text_preview_excel()
    _selftest_generate_text_preview_pdf()
    _selftest_prune_old_previews_keeps_recent_n()
